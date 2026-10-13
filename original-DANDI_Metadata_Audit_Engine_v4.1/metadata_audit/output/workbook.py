from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import AuditConfig, NOT_REPORTED
from ..models import AuditResult, Evidence
from ..processing.mapping import choose_best_evidence
from ..processing.modality import audit_sheet_for_modality

HEADER_FILL = "1F4E78"
MISSING_FILL = "F4CCCC"
REVIEW_FILL = "FFF2CC"
SUCCESS_FILL = "D9EAD3"

CORE_LONG_SHEETS = {
    "Source Records": "Dandiset ID",
    "Publication Candidates": "Dandiset ID",
    "NWB File Metadata": "dandiset_id",
    "Usage Summary": "Dandiset ID",
    "Dataset Features": "Dandiset ID",
}

OBSOLETE_SHEETS = {
    "Dashboard", "Asset Inventory", "NWB Objects", "Electrode Groups",
    "Publication Records", "Web Page Records", "Discovery Records", "GitHub Records",
    "Usage Daily", "Usage Regions", "Usage Assets", "Usage Asset Types",
    "Feature Evidence", "Feature–Usage Analysis", "Audit Log", "DANDI Metadata",
}

FEATURE_FIELDS = [
    "acquisition_equipment",
    "equipment_manufacturer",
    "acquisition_software",
    "analysis_software",
    "electrode_or_probe",
    "probe_model",
    "stimulation_equipment",
    "experimental_model",
    "injection_materials",
    "titer",
    "excitation_wavelength",
    "emission_wavelength",
    "laser_wavelength",
    "core_diameter",
    "numerical_aperture",
    "terminal_anaesthesia",
]


def _style_header(ws: Any) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit(ws: Any, max_width: int = 44) -> None:
    for col in range(1, ws.max_column + 1):
        width = 12
        for row in range(1, min(ws.max_row, 250) + 1):
            width = max(width, len(str(ws.cell(row, col).value or "")) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(width, max_width)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _records(ws: Any) -> list[dict[str, Any]]:
    if ws.max_row < 2:
        return []
    headers = [str(c.value or "").strip() for c in ws[1]]
    if not any(headers):
        return []
    out = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    return out


def _replace_dataset_records(workbook: Any, sheet_name: str, rows: list[dict[str, Any]], key: str, did: str) -> None:
    existing = _records(workbook[sheet_name]) if sheet_name in workbook.sheetnames else []
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    existing = [r for r in existing if str(r.get(key, "")).zfill(6) != did]
    combined = existing + rows
    ws = workbook.create_sheet(sheet_name)
    if not combined:
        ws["A1"] = "No records"
        return
    columns: list[str] = []
    for row in combined:
        for col in row:
            if col not in columns:
                columns.append(col)
    for idx, col in enumerate(columns, 1):
        ws.cell(1, idx, col)
    _style_header(ws)
    for r_idx, record in enumerate(combined, 2):
        for c_idx, col in enumerate(columns, 1):
            value = record.get(col, "")
            ws.cell(r_idx, c_idx, "" if value is None else value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)


def _ensure_master(config: AuditConfig) -> None:
    config.ensure_directories()
    if config.master_workbook.exists():
        return
    if not config.template_path.exists():
        raise FileNotFoundError(f"Master template not found: {config.template_path.resolve()}")
    shutil.copy2(config.template_path, config.master_workbook)


def _backup(config: AuditConfig) -> None:
    if not config.keep_backup or not config.master_workbook.exists():
        return
    archive = config.output_directory / "archive"
    archive.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(config.master_workbook, archive / f"Metadata_Audit_Master_{stamp}.xlsx")


def _find_header_row(ws: Any) -> tuple[int, int]:
    for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row)):
        for cell in row:
            if str(cell.value or "").strip() == "Dandiset ID":
                return cell.row, cell.column
    raise ValueError(f'Could not find "Dandiset ID" in {ws.title!r}.')


def _find_or_add_row(ws: Any, did: str) -> tuple[int, int, list[str]]:
    header_row, id_col = _find_header_row(ws)
    headers = [str(ws.cell(header_row, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    for row in range(header_row + 1, ws.max_row + 1):
        if str(ws.cell(row, id_col).value or "").strip().zfill(6) == did:
            return header_row, row, headers
    row = ws.max_row + 1
    ws.cell(row, id_col, did)
    return header_row, row, headers


def _remove_from_other_audits(workbook: Any, did: str, keep: str) -> None:
    for name in ("Ephys Audit", "Fiber Photometry Audit", "Calcium Imaging Audit", "Other Modalities Audit"):
        if name == keep or name not in workbook.sheetnames:
            continue
        ws = workbook[name]
        try:
            header_row, id_col = _find_header_row(ws)
        except ValueError:
            continue
        for row in range(ws.max_row, header_row, -1):
            if str(ws.cell(row, id_col).value or "").strip().zfill(6) == did:
                ws.delete_rows(row, 1)


def _bytes_human(total: Any) -> str:
    try:
        value = float(total)
    except (TypeError, ValueError):
        return ""
    units = ["B", "KB", "MB", "GB", "TB"]
    idx = 0
    while value >= 1000 and idx < len(units) - 1:
        value /= 1000
        idx += 1
    return f"{value:.2f} {units[idx]}"


def _dataset_size(result: AuditResult) -> str:
    total = 0.0
    found = False
    for row in result.asset_rows:
        try:
            total += float(row.get("asset_size_bytes") or 0)
            found = True
        except (TypeError, ValueError):
            continue
    return _bytes_human(total) if found else ""


def _populate_audit(workbook: Any, config: AuditConfig, result: AuditResult) -> tuple[str, list[str], dict[str, Evidence]]:
    classification = result.modality_classification
    sheet_name = config.audit_sheet if config.audit_sheet != "auto" else audit_sheet_for_modality(classification.get("primary_modality", ""))
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Audit sheet not found: {sheet_name}")
    _remove_from_other_audits(workbook, config.dandiset_id, sheet_name)

    ws = workbook[sheet_name]
    header_row, row, headers = _find_or_add_row(ws, config.dandiset_id)
    metadata_headers = [h for h in headers[9:] if h]
    selected = choose_best_evidence(metadata_headers, result.evidence)

    catalyst = result.catalyst_record
    dois = "; ".join(str(r.get("DOI", "")) for r in result.publication_rows if r.get("DOI"))
    fixed = {
        "Dandiset ID": config.dandiset_id,
        "Contributor": catalyst.get("PI", ""),
        "Brain Region": catalyst.get("Brain Region", ""),
        "Species": catalyst.get("Species", ""),
        "Size": _dataset_size(result),
        "Catalyst Source": catalyst.get("Source", ""),
        "Publication DOI": dois,
        "Audit Status": "Complete" if config.max_assets is None else f"Pilot ({config.max_assets} NWB files max)",
        "Auditor Notes": (
            f"Catalyst modality: {classification.get('catalyst_modality') or 'Not available'}; "
            f"NWB/DANDI detected: {'; '.join(classification.get('detected_modalities', [])) or 'Unclassified'}; "
            f"agreement: {classification.get('agreement', '')}. "
            "Yellow cells are narrative/medium-confidence candidates and require manual confirmation."
        ),
    }

    for col, header in enumerate(headers, 1):
        if not header:
            continue
        if header in fixed:
            ws.cell(row, col, fixed[header])
        elif col > 9:
            ev = selected[header]
            ws.cell(row, col, ev.value)
        cell = ws.cell(row, col)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if col > 9:
            ev = selected.get(header)
            if cell.value == NOT_REPORTED:
                cell.fill = PatternFill("solid", fgColor=MISSING_FILL)
            elif ev and ev.review_status in {"Needs Review", "Needs Manual Review"}:
                cell.fill = PatternFill("solid", fgColor=REVIEW_FILL)
            else:
                cell.fill = PatternFill("solid", fgColor=SUCCESS_FILL)
    ws.freeze_panes = f"A{header_row + 1}"
    return sheet_name, metadata_headers, selected


def _ephys_subtype(catalyst_modality: str) -> str:
    text = str(catalyst_modality or "").lower()
    has_extra = "extra ephys" in text or "extracellular" in text
    has_intra = "intra ephys" in text or "intracellular" in text
    if has_extra and has_intra:
        return "Mixed extracellular/intracellular"
    if has_extra:
        return "Extracellular electrophysiology"
    if has_intra:
        return "Intracellular electrophysiology"
    return ""


def _upsert_inventory(workbook: Any, config: AuditConfig, result: AuditResult, audit_sheet: str) -> None:
    required = [
        "Dandiset ID", "PI", "Catalyst Modality", "Ephys Subtype", "Detected NWB Modalities", "Modality Agreement",
        "Species", "Brain Region", "Year", "Dataset Status", "DANDI Version", "Title",
        "Audit Sheet", "Audit Status", "Last Audited", "NWB Files Read", "NWB Files Failed",
    ]
    old_records = _records(workbook["Dataset Inventory"]) if "Dataset Inventory" in workbook.sheetnames else []
    if "Dataset Inventory" in workbook.sheetnames:
        del workbook["Dataset Inventory"]
    # Preserve only already-upgraded rows; old v1.3 'Modality=Ephys' rows are intentionally discarded.
    preserved = [r for r in old_records if "Catalyst Modality" in r and str(r.get("Dandiset ID", "")).zfill(6) != config.dandiset_id]

    successful = sum(r.get("extraction_status") == "Success" for r in result.nwb_rows)
    failed = sum(r.get("extraction_status") == "Failed" for r in result.nwb_rows)
    cat = result.catalyst_record
    cls = result.modality_classification
    current = {
        "Dandiset ID": config.dandiset_id,
        "PI": cat.get("PI", ""),
        "Catalyst Modality": cat.get("Catalyst Modality", ""),
        "Ephys Subtype": _ephys_subtype(cat.get("Catalyst Modality", "")),
        "Detected NWB Modalities": "; ".join(cls.get("detected_modalities", [])),
        "Modality Agreement": cls.get("agreement", ""),
        "Species": cat.get("Species", ""),
        "Brain Region": cat.get("Brain Region", ""),
        "Year": cat.get("Year", ""),
        "Dataset Status": cat.get("Dataset Status", ""),
        "DANDI Version": config.version,
        "Title": result.dandiset_metadata.get("name", ""),
        "Audit Sheet": audit_sheet,
        "Audit Status": "Complete" if config.max_assets is None else "Pilot",
        "Last Audited": datetime.now().isoformat(timespec="seconds"),
        "NWB Files Read": successful,
        "NWB Files Failed": failed,
    }
    rows = preserved + [current]
    ws = workbook.create_sheet("Dataset Inventory", 1)
    for c, h in enumerate(required, 1):
        ws.cell(1, c, h)
    _style_header(ws)
    for r_idx, record in enumerate(rows, 2):
        for c_idx, h in enumerate(required, 1):
            ws.cell(r_idx, c_idx, record.get(h, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)


def _dataset_features(config: AuditConfig, result: AuditResult) -> list[dict[str, Any]]:
    by_field: dict[str, list[Evidence]] = {}
    for ev in result.evidence:
        if ev.audit_field in FEATURE_FIELDS:
            by_field.setdefault(ev.audit_field, []).append(ev)

    record: dict[str, Any] = {
        "Dandiset ID": config.dandiset_id,
        "PI": result.catalyst_record.get("PI", ""),
        "Catalyst Modality": result.catalyst_record.get("Catalyst Modality", ""),
        "Ephys Subtype": _ephys_subtype(result.catalyst_record.get("Catalyst Modality", "")),
        "Brain Region": result.catalyst_record.get("Brain Region", ""),
    }

    confidence_rank = {"High": 3, "Medium": 2, "Low": 1, "—": 0}
    for field in FEATURE_FIELDS:
        candidates = by_field.get(field, [])
        candidates = sorted(
            candidates,
            key=lambda ev: confidence_rank.get(ev.confidence, 0),
            reverse=True,
        )
        values: list[str] = []
        sources: list[str] = []
        confidences: list[str] = []
        statuses: list[str] = []
        for ev in candidates:
            value = str(ev.value)
            if value not in values:
                values.append(value)
            if ev.source_type not in sources:
                sources.append(ev.source_type)
            if ev.confidence not in confidences:
                confidences.append(ev.confidence)
            if ev.review_status not in statuses:
                statuses.append(ev.review_status)

        record[field] = "; ".join(values[:6]) if values else NOT_REPORTED
        record[f"{field} source"] = "; ".join(sources) if sources else ""
        record[f"{field} confidence"] = "; ".join(confidences) if confidences else ""
        record[f"{field} review"] = "; ".join(statuses) if statuses else ""

    narrative_count = sum(
        1 for ev in result.evidence
        if ev.audit_field in FEATURE_FIELDS and ev.review_status == "Needs Manual Review"
    )
    record["Review Status"] = (
        f"{narrative_count} narrative candidate(s) need manual review"
        if narrative_count
        else "No unresolved narrative feature candidates"
    )
    return [record]


def _write_cde(workbook: Any, config: AuditConfig, audit_sheet: str, metadata_headers: list[str], selected: dict[str, Evidence]) -> None:
    existing = _records(workbook["Common Data Elements"]) if "Common Data Elements" in workbook.sheetnames else []
    # If this is the static field dictionary from the template, replace it with actual audit coverage.
    existing = [r for r in existing if "Dandiset ID" in r and str(r.get("Dandiset ID", "")).zfill(6) != config.dandiset_id]
    rows = existing + [{
        "Dandiset ID": config.dandiset_id,
        "Audit Sheet": audit_sheet,
        "Common Data Element": header,
        "Value": selected[header].value,
        "Present": "No" if selected[header].value == NOT_REPORTED else "Yes",
        "Best Source": selected[header].source_type,
        "Confidence": selected[header].confidence,
        "Review Status": selected[header].review_status,
    } for header in metadata_headers]
    if "Common Data Elements" in workbook.sheetnames:
        del workbook["Common Data Elements"]
    ws = workbook.create_sheet("Common Data Elements")
    columns = ["Dandiset ID", "Audit Sheet", "Common Data Element", "Value", "Present", "Best Source", "Confidence", "Review Status"]
    for c, h in enumerate(columns, 1):
        ws.cell(1, c, h)
    _style_header(ws)
    for r_idx, rec in enumerate(rows, 2):
        for c_idx, h in enumerate(columns, 1):
            ws.cell(r_idx, c_idx, rec.get(h, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)


def write_master_workbook(config: AuditConfig, result: AuditResult) -> Path:
    _ensure_master(config)
    _backup(config)
    try:
        workbook = load_workbook(config.master_workbook)
    except PermissionError as exc:
        raise PermissionError(f"Close {config.master_workbook.name} in Excel and run again.") from exc

    for sheet in list(OBSOLETE_SHEETS):
        if sheet in workbook.sheetnames:
            del workbook[sheet]

    audit_sheet, metadata_headers, selected = _populate_audit(workbook, config, result)
    _upsert_inventory(workbook, config, result, audit_sheet)

    source_rows = [ev.as_dict(config.dandiset_id) for ev in result.evidence]
    publication_rows = [r for r in source_rows if r.get("Source Type") == "Confirmed publication methods (PMC)"]
    datasets = {
        "Source Records": source_rows,
        "Publication Candidates": publication_rows,
        "NWB File Metadata": result.nwb_rows,
        "Usage Summary": result.usage_summary,
        "Dataset Features": _dataset_features(config, result),
    }
    for sheet, rows in datasets.items():
        _replace_dataset_records(workbook, sheet, rows, CORE_LONG_SHEETS[sheet], config.dandiset_id)

    _write_cde(workbook, config, audit_sheet, metadata_headers, selected)

    # Put the most useful sheets first.
    preferred = [
        "README", "Dataset Inventory", "Ephys Audit", "Fiber Photometry Audit", "Calcium Imaging Audit",
        "Other Modalities Audit", "Dataset Features", "Common Data Elements", "Source Records",
        "Publication Candidates", "NWB File Metadata", "Usage Summary",
    ]
    ordered = [workbook[name] for name in preferred if name in workbook.sheetnames]
    ordered += [ws for ws in workbook.worksheets if ws.title not in preferred]
    workbook._sheets = ordered

    temp = config.master_workbook.with_name(f".{config.master_workbook.stem}.saving{config.master_workbook.suffix}")
    try:
        workbook.save(temp)
        temp.replace(config.master_workbook)
    except PermissionError as exc:
        temp.unlink(missing_ok=True)
        raise PermissionError(f"Close {config.master_workbook.name} in Excel and run again.") from exc
    return config.master_workbook
