from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from metadata_audit.config import AuditConfig
from metadata_audit.models import AuditResult, Evidence
from metadata_audit.output.workbook import write_master_workbook


ROOT = Path(__file__).resolve().parents[1]


def make_result(status: str) -> AuditResult:
    return AuditResult(
        dandiset_metadata={"name": "Example", "description": "Dataset-level context"},
        dandi_reference_rows=[{
            "Dandiset ID": "001177", "Asset ID": "", "Audit Field": "name",
            "Extracted Value": "Mus musculus", "Source Type": "DANDI landing page",
            "Source Location": "assetsSummary.species[0].name",
            "Review Status": "Reference only", "Evidence Role": "DANDI reference",
        }],
        nwb_rows=[{"dandiset_id": "001177", "asset_path": "sub-1.nwb",
                   "extraction_status": status}],
        modality_classification={"primary_modality": "Intracellular electrophysiology",
                                 "detected_modalities": ["Intracellular electrophysiology"] if status == "Success" else []},
        evidence=[Evidence(audit_field="sex", value="F", source_type="NWB file",
                           source_location="sub-1.nwb:/general/subject/sex",
                           review_status="NWB structured value", confidence="High")]
                 if status == "Success" else [],
    )


class SourceSeparationTest(unittest.TestCase):
    def generate(self, status):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            config = AuditConfig("001177", master_workbook=directory / "result.xlsx",
                                 template_path=ROOT / "templates/Metadata_Audit_Master_Template.xlsx",
                                 output_directory=directory, cache_directory=directory / "cache",
                                 max_assets=1)
            result = make_result(status)
            book = load_workbook(write_master_workbook(config, result), data_only=True)
            head = [cell.value for cell in book["Common Data Elements"][1]]
            data = [dict(zip(head, row)) for row in book["Common Data Elements"].iter_rows(min_row=2, values_only=True)]
            sources = list(book["Source Records"].values)
            yield book, data, sources

    def test_dandi_reference_does_not_fill_nwb_value(self):
        for book, data, sources in self.generate("Success"):
            species = next(row for row in data if row["Common Data Element"] == "species")
            self.assertEqual(species["DANDI Reference (not NWB)"], "Mus musculus")
            self.assertFalse(species["NWB Value"])
            self.assertNotEqual(species["Present in NWB"], "Yes")
            self.assertTrue(any("DANDI reference" in row for row in sources[1:]))
            self.assertEqual(book["Source Records"]["B1"].value, "Asset ID")
            self.assertEqual(book["Common Data Elements"]["B1"].value, "Field ID")

    def test_failed_read_is_unknown(self):
        for _, data, _ in self.generate("Failed"):
            species = next(row for row in data if row["Common Data Element"] == "species")
            self.assertEqual(species["Present in NWB"], "Could not assess")


if __name__ == "__main__":
    unittest.main()
