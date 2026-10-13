from __future__ import annotations

import shutil
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import AuditConfig, NOT_REPORTED
from ..models import AuditResult, Evidence
from ..processing.mapping import CONTEXTUAL_FIELDS, choose_best_evidence, matches_context, modality_observed
from ..processing.modality import audit_sheet_for_modality

HEADER_FILL = "75405C"
MISSING_FILL = "F4EFF2"
REVIEW_FILL = "FFF1D9"
SUCCESS_FILL = "E9F2EB"
STRIPE_FILL = "F4E9EF"
BASE_FILL = "FCF7FA"
REQUIREMENT_COLORS = {
    "Required": "F4CCCC",
    "Highly Recommended": "FFF2CC",
    "Conditionally Required": "D9EAD3",
}

CORE_LONG_SHEETS = {
    "Requirement Review": "Dandiset ID",
    "NWB Properties": "dandiset_id",
    "DANDI Resources": "Dandiset ID",
    "Source Records": "Dandiset ID",
    "NWB File Metadata": "dandiset_id",
    "Dataset Features": "Dandiset ID",
    "NWB Findings": "Dandiset ID",
    "NWB Objects": "dandiset_id",
    "Electrode Groups": "dandiset_id",
    "Asset Inventory": "dandiset_id",
}

OBSOLETE_SHEETS = {
    "Publication Candidates", "Usage Summary",
    "Dashboard",
    "Publication Records", "Web Page Records", "Discovery Records", "GitHub Records",
    "Usage Daily", "Usage Regions", "Usage Assets", "Usage Asset Types",
    "Feature Evidence", "Feature–Usage Analysis", "Audit Log", "DANDI Metadata",
}

FEATURE_FIELDS = [
    "acquisition_equipment",
    "equipment_manufacturer",
    "acquisition_software",
    "analysis_software",
    "electrode_or_probe",
    "probe_model",
    "stimulation_equipment",
    "experimental_model",
    "injection_materials",
    "titer",
    "excitation_wavelength",
    "emission_wavelength",
    "laser_wavelength",
    "core_diameter",
    "numerical_aperture",
    "terminal_anaesthesia",
]


def _style_header(ws: Any) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit(ws: Any, max_width: int = 44) -> None:
    for col in range(1, ws.max_column + 1):
        width = 12
        for row in range(1, min(ws.max_row, 250) + 1):
            width = max(width, len(str(ws.cell(row, col).value or "")) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(width, max_width)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _band_rows(ws: Any) -> None:
    if ws.max_row > 1:
        headers = [str(cell.value or "") for cell in ws[1]]
        req_col = next((i for i, name in enumerate(headers, 1) if name in {"Requirement Level", "Requirement level"}), None)
        status_col = headers.index("Present in NWB") + 1 if "Present in NWB" in headers else None
        for row in ws.iter_rows(min_row=2):
            tint = STRIPE_FILL if row[0].row % 2 == 0 else BASE_FILL
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=tint)
            if req_col and str(row[req_col-1].value) in REQUIREMENT_COLORS:
                row[req_col-1].fill = PatternFill("solid", fgColor=REQUIREMENT_COLORS[str(row[req_col-1].value)])
            if status_col:
                state = str(row[status_col-1].value or "")
                color = SUCCESS_FILL if state == "Yes" else REVIEW_FILL if state in {"To review", "Inspection incomplete", "Review context"} else MISSING_FILL
                row[status_col-1].fill = PatternFill("solid", fgColor=color)


def _records(ws: Any) -> list[dict[str, Any]]:
    if ws.max_row < 2:
        return []
    headers = [str(c.value or "").strip() for c in ws[1]]
    if not any(headers):
        return []
    out = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        out.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    return out


def _replace_dataset_records(workbook: Any, sheet_name: str, rows: list[dict[str, Any]], key: str, did: str) -> None:
    existing = _records(workbook[sheet_name]) if sheet_name in workbook.sheetnames else []
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    existing = [r for r in existing if str(r.get(key, "")).zfill(6) != did]
    combined = existing + rows
    ws = workbook.create_sheet(sheet_name)
    if not combined:
        fallback = {
            "NWB Properties": ["dandiset_id", "asset_id", "asset_path", "object_path", "property", "value", "read_status"],
            "NWB Objects": ["dandiset_id", "asset_id", "asset_path", "object_category", "object_name", "neurodata_type"],
            "Electrode Groups": ["dandiset_id", "asset_id", "asset_path", "electrode_group_name", "location", "device"],
            "Asset Inventory": ["dandiset_id", "version", "asset_id", "asset_path"],
            "Source Records": ["Dandiset ID", "Asset ID", "Audit Field", "Extracted Value", "Source Type", "Source Location", "Evidence Role"],
        }
        headers = fallback.get(sheet_name, [key])
        for idx, header in enumerate(headers, 1):
            ws.cell(1, idx, header)
        _style_header(ws)
        return
    columns: list[str] = []
    for row in combined:
        for col in row:
            if col not in columns:
                columns.append(col)
    for idx, col in enumerate(columns, 1):
        ws.cell(1, idx, col)
    _style_header(ws)
    for r_idx, record in enumerate(combined, 2):
        for c_idx, col in enumerate(columns, 1):
            value = record.get(col, "")
            ws.cell(r_idx, c_idx, "" if value is None else value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)
    _band_rows(ws)


def _ensure_master(config: AuditConfig) -> None:
    config.ensure_directories()
    if config.master_workbook.exists():
        return
    if not config.template_path.exists():
        raise FileNotFoundError(f"Master template not found: {config.template_path.resolve()}")
    shutil.copy2(config.template_path, config.master_workbook)


def _backup(config: AuditConfig) -> None:
    if not config.keep_backup or not config.master_workbook.exists():
        return
    archive = config.output_directory / "archive"
    archive.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(config.master_workbook, archive / f"Metadata_Audit_Master_{stamp}.xlsx")


def _find_header_row(ws: Any) -> tuple[int, int]:
    for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row)):
        for cell in row:
            if str(cell.value or "").strip() == "Dandiset ID":
                return cell.row, cell.column
    raise ValueError(f'Could not find "Dandiset ID" in {ws.title!r}.')


def _find_or_add_row(ws: Any, did: str) -> tuple[int, int, list[str]]:
    header_row, id_col = _find_header_row(ws)
    headers = [str(ws.cell(header_row, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    for row in range(header_row + 1, ws.max_row + 1):
        if str(ws.cell(row, id_col).value or "").strip().zfill(6) == did:
            return header_row, row, headers
    row = ws.max_row + 1
    ws.cell(row, id_col, did)
    return header_row, row, headers


def _remove_from_other_audits(workbook: Any, did: str, keep: str) -> None:
    for name in ("Ephys Audit", "Fiber Photometry Audit", "Calcium Imaging Audit", "Other Modalities Audit"):
        if name == keep or name not in workbook.sheetnames:
            continue
        ws = workbook[name]
        try:
            header_row, id_col = _find_header_row(ws)
        except ValueError:
            continue
        for row in range(ws.max_row, header_row, -1):
            if str(ws.cell(row, id_col).value or "").strip().zfill(6) == did:
                ws.delete_rows(row, 1)


def _bytes_human(total: Any) -> str:
    try:
        value = float(total)
    except (TypeError, ValueError):
        return ""
    units = ["B", "KB", "MB", "GB", "TB"]
    idx = 0
    while value >= 1000 and idx < len(units) - 1:
        value /= 1000
        idx += 1
    return f"{value:.2f} {units[idx]}"


def _dataset_size(result: AuditResult) -> str:
    total = 0.0
    found = False
    for row in result.asset_rows:
        try:
            total += float(row.get("asset_size_bytes") or 0)
            found = True
        except (TypeError, ValueError):
            continue
    return _bytes_human(total) if found else ""


def _populate_audit(workbook: Any, config: AuditConfig, result: AuditResult) -> tuple[str, list[str], dict[str, Evidence]]:
    classification = result.modality_classification
    sheet_name = config.audit_sheet if config.audit_sheet != "auto" else audit_sheet_for_modality(classification.get("primary_modality", ""))
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Audit sheet not found: {sheet_name}")
    _remove_from_other_audits(workbook, config.dandiset_id, sheet_name)

    ws = workbook[sheet_name]
    header_row, row, headers = _find_or_add_row(ws, config.dandiset_id)
    metadata_headers = [h for h in headers[8:] if h and h != "Auditor Notes"]
    catalog = json.loads((Path(__file__).resolve().parents[1] / "field_catalog.json").read_text(encoding="utf-8"))
    permitted = {
        entry["field"] for entry in catalog
        if entry["modality"] in ({"Shared", "Electrophysiology"} if sheet_name == "Ephys Audit" else
                                 {"Shared", "Fiber photometry"} if sheet_name == "Fiber Photometry Audit" else
                                 {"Shared", "Calcium imaging"} if sheet_name == "Calcium Imaging Audit" else {"Shared"})
        and modality_observed(entry["modality"], entry["experiment_context"],
                              set(classification.get("detected_modalities", [])))
    }
    selected = choose_best_evidence(metadata_headers, result.evidence, permitted)

    fixed = {
        "Dandiset ID": config.dandiset_id,
        "Contributor": "",
        "Brain Region": "",
        "Species": "; ".join(dict.fromkeys(str(r["species"]) for r in result.nwb_rows
                                               if r.get("extraction_status") == "Success" and r.get("species"))),
        "Size": "",  # DANDI asset sizes remain in Asset Inventory.
        "Catalyst Source": "",
        "Publication DOI": "",
        "Audit Status": f"Pilot: {len(result.nwb_rows)} of {sum(str(a.get('asset_path','')).lower().endswith('.nwb') for a in result.asset_rows)} NWB files" if config.max_assets is not None else "Review NWB read status" if not result.nwb_rows or any(r.get("extraction_status") != "Success" or r.get("metadata_read_errors") for r in result.nwb_rows) else "NWB inspection complete",
        "Auditor Notes": "See NWB Findings for extracted subject, session, device and recording facts. Blank fields have no validated match in scanned files; they are not proven absent from the Dandiset.",
    }

    for col, header in enumerate(headers, 1):
        if not header:
            continue
        if header in fixed:
            ws.cell(row, col, fixed[header])
        elif header in selected:
            ev = selected[header]
            ws.cell(row, col, "" if ev.value == NOT_REPORTED else ev.value)
        cell = ws.cell(row, col)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor=BASE_FILL)
        if header in selected:
            ev = selected.get(header)
            if not cell.value:
                cell.fill = PatternFill("solid", fgColor=MISSING_FILL)
            elif ev and ev.review_status in {"Needs Review", "Needs Manual Review", "NWB text—review"}:
                cell.fill = PatternFill("solid", fgColor=REVIEW_FILL)
            else:
                cell.fill = PatternFill("solid", fgColor=SUCCESS_FILL)
    ws.freeze_panes = f"A{header_row + 1}"
    return sheet_name, metadata_headers, selected


def _upsert_inventory(workbook: Any, config: AuditConfig, result: AuditResult, audit_sheet: str) -> None:
    required = [
        "Dandiset ID", "DANDI Version", "Title (DANDI reference)", "Contributors (DANDI reference)",
        "Description (DANDI reference)", "Detected NWB Modalities", "Audit Sheet",
        "NWB Files In Dandiset", "NWB Files Read", "NWB Files Failed", "NWB Files Not Scanned",
        "Files With Metadata Errors", "Audit Status", "Last Audited",
        "NWB Properties Preview Rows", "Full NWB Metadata File", "DANDI Resource Count",
    ]
    old_records = _records(workbook["Dataset Inventory"]) if "Dataset Inventory" in workbook.sheetnames else []
    if "Dataset Inventory" in workbook.sheetnames:
        del workbook["Dataset Inventory"]
    preserved = [r for r in old_records if "Title (DANDI reference)" in r
                 and str(r.get("Dandiset ID", "")).zfill(6) != config.dandiset_id]

    successful = sum(r.get("extraction_status") == "Success" for r in result.nwb_rows)
    failed = sum(r.get("extraction_status") == "Failed" for r in result.nwb_rows)
    total_nwb = sum(str(a.get("asset_path", "")).lower().endswith(".nwb") for a in result.asset_rows)
    errors = sum(bool(r.get("metadata_read_errors")) for r in result.nwb_rows)
    cls = result.modality_classification
    contributors = result.dandiset_metadata.get("contributor") or []
    current = {
        "Dandiset ID": config.dandiset_id,
        "DANDI Version": config.version,
        "Title (DANDI reference)": result.dandiset_metadata.get("name", ""),
        "Contributors (DANDI reference)": "; ".join(str(c.get("name", "")) for c in contributors if isinstance(c, dict)),
        "Description (DANDI reference)": str(result.dandiset_metadata.get("description") or "")[:3000],
        "Detected NWB Modalities": "; ".join(cls.get("detected_modalities", [])),
        "Audit Sheet": audit_sheet,
        "Audit Status": "No NWB files found" if not result.nwb_rows else "Review read errors" if failed or errors else "NWB inspected" if config.max_assets is None else "NWB pilot",
        "Last Audited": datetime.now().isoformat(timespec="seconds"),
        "NWB Files Read": successful,
        "NWB Files Failed": failed,
        "NWB Files In Dandiset": total_nwb,
        "NWB Files Not Scanned": max(0, total_nwb - len(result.nwb_rows)),
        "Files With Metadata Errors": errors,
        "NWB Properties Preview Rows": len(result.nwb_property_rows),
        "Full NWB Metadata File": str(config.output_directory / f"NWB_Properties_{config.dandiset_id}.jsonl.gz"),
        "DANDI Resource Count": len(result.associated_resources),
    }
    rows = preserved + [current]
    ws = workbook.create_sheet("Dataset Inventory", 1)
    for c, h in enumerate(required, 1):
        ws.cell(1, c, h)
    _style_header(ws)
    for r_idx, record in enumerate(rows, 2):
        for c_idx, h in enumerate(required, 1):
            ws.cell(r_idx, c_idx, record.get(h, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)
    _band_rows(ws)


def _dataset_features(config: AuditConfig, result: AuditResult) -> list[dict[str, Any]]:
    by_field: dict[str, list[Evidence]] = {}
    for ev in result.evidence:
        if ev.audit_field in FEATURE_FIELDS:
            by_field.setdefault(ev.audit_field, []).append(ev)

    successful = [r for r in result.nwb_rows if r.get("extraction_status") == "Success"]
    def observed(column: str) -> str:
        values = list(dict.fromkeys(str(r[column]) for r in successful if r.get(column)))
        return "; ".join(values[:12])
    record: dict[str, Any] = {
        "Dandiset ID": config.dandiset_id,
        "Detected NWB Modalities": "; ".join(result.modality_classification.get("detected_modalities", [])),
        "NWB Files Read": len(successful),
        "NWB Files Failed": sum(r.get("extraction_status") == "Failed" for r in result.nwb_rows),
        "Subject Species (NWB)": observed("species"),
        "Device Names (NWB)": observed("device_names"),
        "Electrode Locations (NWB)": observed("electrode_locations"),
        "Acquisition Objects (NWB)": observed("acquisition_object_names"),
        "Processing Modules (NWB)": observed("processing_module_names"),
    }

    confidence_rank = {"High": 3, "Medium": 2, "Low": 1, "—": 0}
    for field in FEATURE_FIELDS:
        candidates = by_field.get(field, [])
        if not candidates:
            continue
        candidates = sorted(
            candidates,
            key=lambda ev: confidence_rank.get(ev.confidence, 0),
            reverse=True,
        )
        values: list[str] = []
        sources: list[str] = []
        confidences: list[str] = []
        statuses: list[str] = []
        for ev in candidates:
            value = str(ev.value)
            if value not in values:
                values.append(value)
            if ev.source_type not in sources:
                sources.append(ev.source_type)
            if ev.confidence not in confidences:
                confidences.append(ev.confidence)
            if ev.review_status not in statuses:
                statuses.append(ev.review_status)

        record[field] = "; ".join(values[:6]) if values else ""
        record[f"{field} source"] = "; ".join(sources) if sources else ""
        record[f"{field} confidence"] = "; ".join(confidences) if confidences else ""
        record[f"{field} review"] = "; ".join(statuses) if statuses else ""

    record["Review Status"] = "Review generic field mappings and file failures before reporting completeness"
    return [record]


def _nwb_findings(config: AuditConfig, result: AuditResult) -> list[dict[str, Any]]:
    """Small, sourced per-file view; raw tables and arrays stay in NWB Properties."""
    from collections import defaultdict
    output: list[dict[str, Any]] = []
    fields = {
        "subject_id": ("Subject", "Subject ID"), "sex": ("Subject", "Sex"),
        "date_of_birth": ("Subject", "Date of birth"), "species": ("Subject", "Species"),
        "genotype": ("Subject", "Genotype"), "subject_description": ("Subject", "Subject description"),
        "session_id": ("Session", "Session ID"), "session_start_time": ("Session", "Session start"),
        "session_description": ("Session", "Session description"),
        "experimenter": ("Session", "Experimenter"), "institution": ("Session", "Institution"),
        "lab": ("Session", "Lab"), "device_names": ("Recording", "Device names"),
        "electrode_count": ("Recording", "Electrode count"),
        "electrode_locations": ("Recording", "Electrode group location"),
        "processing_module_names": ("Recording", "Processing modules"),
    }
    for file in result.nwb_rows:
        if file.get("extraction_status") != "Success":
            continue
        asset = file.get("asset_path", "")
        aid = file.get("asset_id", "")
        props = [p for p in result.nwb_property_rows if p.get("asset_id") == aid]
        for key, (section, label) in fields.items():
            value = file.get(key)
            if value is None or str(value).strip() == "":
                continue
            source = next((p for p in props if p.get("property") in (key, "[scalar]")
                           and (p.get("property") == key or str(p.get("object_path", "")).endswith("/" + key))
                           and str(p.get("object_path", "")).startswith("/")), None)
            fallback = {"subject_description": "/general/subject/description",
                        "device_names": "/general/devices", "electrode_count": "/general/extracellular_ephys/electrodes",
                        "electrode_locations": "/general/extracellular_ephys/*/location",
                        "processing_module_names": "/processing"}
            path = str(source["object_path"]) if source and source.get("property") == "[scalar]" else \
                   str(source["object_path"]) + "/" + key if source else fallback.get(key, "NWB File Metadata:" + key)
            narrative = key.endswith("description")
            interpretation = ("Narrative in NWB; review context" if narrative else
                              "Value stored as unknown" if str(value).lower() == "unknown" else "Direct NWB metadata")
            if any(token in str(value) for token in ("â€", "Ã", "�")):
                interpretation += "; possible text encoding issue"
            output.append({"Dandiset ID": config.dandiset_id, "Asset path": asset,
                           "Section": section, "NWB finding": label,
                           "Extracted value": str(value)[:2000], "NWB path": path,
                           "Interpretation": interpretation})
        grouped: dict[tuple[str, str, str, str], list[str]] = defaultdict(list)
        for obj in result.nwb_object_rows:
            if obj.get("asset_id") != aid or obj.get("object_category") not in ("acquisition", "stimulus", "processing_module", "nested"):
                continue
            if obj.get("object_category") == "nested" and obj.get("neurodata_type") not in {"ElectricalSeries", "RoiResponseSeries", "TwoPhotonSeries", "OnePhotonSeries"}:
                continue
            if obj.get("object_category") == "nested" and any(
                candidate.get("asset_id") == aid and candidate.get("object_category") == "acquisition"
                and candidate.get("object_name") == obj.get("object_name")
                and candidate.get("rate") == obj.get("rate") for candidate in result.nwb_object_rows):
                continue
            key = (str(obj.get("object_category")), str(obj.get("neurodata_type")),
                   str(obj.get("rate") or ""), str(obj.get("unit") or ""))
            grouped[key].append(str(obj.get("object_name", "")))
        for (category, kind, rate, unit), names in grouped.items():
            value = f'{len(names)} object(s): {", ".join(names[:4])}'
            if len(names) > 4:
                value += f' (and {len(names) - 4} more)'
            if rate:
                value += f'; rate {rate} Hz'
            if unit:
                value += f'; unit {unit}'
            output.append({"Dandiset ID": config.dandiset_id, "Asset path": asset,
                           "Section": "NWB objects", "NWB finding": f"{category}: {kind}",
                           "Extracted value": value,
                           "NWB path": "/processing (nested object; exact path in NWB Properties)" if category == "nested"
                                       else "/processing" if category == "processing_module" else f"/{category}",
                           "Interpretation": "Object summary; see NWB Objects for each item"})
        table_cells: dict[str, list[str]] = defaultdict(list)
        for prop in props:
            if (prop.get("evidence_kind") == "NWB table cell" and
                    str(prop.get("object_path", "")).endswith("electrodes")):
                table_cells[str(prop.get("property"))].append(str(prop.get("value", "")))
        for name, values in table_cells.items():
            values = [v for v in values if v]
            if not values:
                continue
            unique = list(dict.fromkeys(values))
            display = "; ".join(unique[:6])
            if len(unique) > 6:
                display += f"; … ({len(unique)} distinct values)"
            output.append({"Dandiset ID": config.dandiset_id, "Asset path": asset,
                           "Section": "Electrode table", "NWB finding": name,
                           "Extracted value": display, "NWB path": f"/general/extracellular_ephys/electrodes/{name}",
                           "Interpretation": f"{len(values)} rows; see NWB Properties for individual cells"})
    return output


def _requirement_review(config: AuditConfig, result: AuditResult) -> list[dict[str, Any]]:
    catalog_path = Path(__file__).resolve().parents[1] / "field_catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    detected = set(result.modality_classification.get("detected_modalities", []))
    detected_text = " ".join(detected).lower()
    successful = sum(r.get("extraction_status") == "Success" for r in result.nwb_rows)
    failed = sum(r.get("extraction_status") == "Failed" for r in result.nwb_rows)
    errors = sum(bool(r.get("metadata_read_errors")) for r in result.nwb_rows)
    by_field: dict[str, list[Evidence]] = {}
    for evidence in result.evidence:
        by_field.setdefault(evidence.audit_field.casefold(), []).append(evidence)
    rows = []
    def dandi_reference(field: str) -> tuple[str, str]:
        # Explicit contextual references only; a generic DANDI `name` cannot
        # establish a viral-material, probe, or channel name.
        if field != "species":
            return "", ""
        candidates = [r for r in result.dandi_reference_rows
                      if "assetsSummary.species[" in r.get("Source Location", "")
                      and r.get("Source Location", "").endswith(".name")]
        return ("; ".join(dict.fromkeys(r["Extracted Value"] for r in candidates)),
                "; ".join(dict.fromkeys(r["Source Location"] for r in candidates)))
    for entry in catalog:
        field = entry["field"]
        modality = entry["modality"]
        context = entry["experiment_context"]
        # Keep the entire register visible, including methods not detected in a pilot.
        observed = modality_observed(modality, context, detected)
        contextual = field.casefold() in CONTEXTUAL_FIELDS
        matches = [] if contextual or not observed else [
            ev for ev in by_field.get(field.casefold(), [])
            if matches_context(field, ev.source_location, context)]
        if modality == "Shared":
            applicability = "Applies" if context == "All recorded subjects" else "Review procedure/context"
        elif modality == "Cross-modality":
            applicability = "Review additional method"
        elif modality == "Electrophysiology":
            applicability = "Review experiment subtype" if any("electrophysiology" in d.lower() for d in detected) else "No ephys detected; review"
        elif modality.lower() in " ".join(detected).lower():
            applicability = "Review experiment subtype" if any(x in context.lower() for x in ("ex vivo", "in vivo", "surgery", "implant")) else "Modality detected"
        else:
            applicability = "Modality not detected; review"
        if successful == 0:
            status = "Could not assess—no NWB files read"
        elif not observed:
            status = "Method not detected in scanned NWB files; not assessed"
        elif contextual:
            status = "Contextual mapping required; inspect NWB paths"
        elif matches:
            status = "NWB value found—mapping review" if any(e.review_status != "NWB structured value" for e in matches) else "NWB value found"
        else:
            status = "No matched value in scanned NWB files; check context and other assets" if failed == 0 and errors == 0 else "No matched value; inspection incomplete"
        ref_value, ref_path = dandi_reference(field)
        rows.append({"Dandiset ID": config.dandiset_id, "Field ID": entry["field_id"],
                     "Origin": entry["source"], "Modality": modality,
                     "Experiment context": context, "Evaluate when": entry["applies_when"],
                     "Field": field, "Requirement level": entry["requirement_level"],
                     "Applicability": applicability, "NWB value": "; ".join(dict.fromkeys(str(e.value) for e in matches))[:2000],
                     "NWB paths": "; ".join(dict.fromkeys(e.source_location for e in matches))[:4000],
                     "DANDI reference": ref_value if not matches else "",
                     "DANDI source path": ref_path if not matches else "",
                     "Match status": status, "Files read": successful, "Files failed": failed,
                     "Files with metadata errors": errors,
                     "Source cell": entry["source_reference"]})
    return rows


def _write_cde(workbook: Any, config: AuditConfig, requirement_rows: list[dict[str, Any]]) -> None:
    existing = _records(workbook["Common Data Elements"]) if "Common Data Elements" in workbook.sheetnames else []
    # If this is the static field dictionary from the template, replace it with actual audit coverage.
    existing = [r for r in existing if "Dandiset ID" in r and str(r.get("Dandiset ID", "")).zfill(6) != config.dandiset_id]
    rows = existing + [{
        "Dandiset ID": config.dandiset_id,
        "Field ID": entry["Field ID"],
        "Common Data Element": entry["Field"],
        "Origin": entry["Origin"],
        "Modality": entry["Modality"],
        "Experiment Context": entry["Experiment context"],
        "Requirement Level": entry["Requirement level"],
        "Applicability": entry["Applicability"],
        "NWB Value": entry["NWB value"],
        "Present in NWB": "To review" if entry["NWB value"] and "review" in entry["Match status"].lower() else "Yes" if entry["NWB value"] else "Could not assess" if not entry["Files read"] else "Review context" if entry["Applicability"].startswith("Review") else "Not found in scanned files" if not entry["Files failed"] and not entry["Files with metadata errors"] else "Inspection incomplete",
        "NWB Path": entry["NWB paths"],
        "DANDI Reference (not NWB)": entry["DANDI reference"],
        "DANDI Source Path": entry["DANDI source path"],
        "Review Status": entry["Match status"],
        "NWB Files Read": entry["Files read"],
        "NWB Files Failed": entry["Files failed"],
    } for entry in requirement_rows if "Method not detected" not in entry["Match status"]]
    if "Common Data Elements" in workbook.sheetnames:
        del workbook["Common Data Elements"]
    ws = workbook.create_sheet("Common Data Elements")
    columns = ["Dandiset ID", "Field ID", "Common Data Element", "Origin", "Modality", "Experiment Context",
               "Requirement Level", "Applicability", "NWB Value", "Present in NWB", "NWB Path",
               "DANDI Reference (not NWB)", "DANDI Source Path", "Review Status", "NWB Files Read", "NWB Files Failed"]
    for c, h in enumerate(columns, 1):
        ws.cell(1, c, h)
    _style_header(ws)
    for r_idx, rec in enumerate(rows, 2):
        for c_idx, h in enumerate(columns, 1):
            ws.cell(r_idx, c_idx, rec.get(h, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autofit(ws)
    _band_rows(ws)


def _write_requirements(workbook: Any) -> None:
    """The Allen register is only fields present in the supplied Allen workbook."""
    catalog_path = Path(__file__).resolve().parents[1] / "field_catalog.json"
    rows = json.loads(catalog_path.read_text(encoding="utf-8"))
    columns = ["Field ID", "Origin", "Modality", "Experiment context", "Evaluate when",
               "Section", "Field", "Requirement level", "Definition", "Source cell"]
    for sheet_name, subset in (("Allen Requirements", [r for r in rows if r["source"] == "Allen Institute"]),
                               ("Field Catalog", rows)):
        if sheet_name in workbook:
            del workbook[sheet_name]
        ws = workbook.create_sheet(sheet_name)
        for c, header in enumerate(columns, 1):
            ws.cell(1, c, header)
        for index, row in enumerate(subset, 2):
            values = [row["field_id"], row["source"], row["modality"], row["experiment_context"],
                      row["applies_when"], row["section"], row["field"], row["requirement_level"],
                      row["definition"], row["source_reference"]]
            for c, value in enumerate(values, 1):
                ws.cell(index, c, value)
        _style_header(ws)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        _autofit(ws)
        _band_rows(ws)


def _source_records(config: AuditConfig, result: AuditResult) -> list[dict[str, Any]]:
    source_rows = []
    curated: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for prop in result.nwb_property_rows:
        if prop.get("evidence_kind") not in {"NWB structured field", "NWB table cell", "NWB scalar dataset"}:
            continue  # Technical attributes and arrays remain searchable in NWB Properties/JSONL.
        name = str(prop.get("property", ""))
        value = prop.get("value", "")
        if value in (None, ""):
            continue
        path = str(prop.get("object_path", ""))
        position = f'[row {prop["row_index"]}]' if prop.get("row_index") is not None else ""
        key = (str(prop.get("asset_id", "")), name, str(value),
               str(prop.get("neurodata_type", "")))
        rec = curated.get(key)
        if rec is not None:
            rec["Occurrences"] += 1
            if path.startswith("/") and ":/" not in rec["Source Location"]:
                rec["Source Location"] = f'{prop.get("asset_path", "")}:{path}{"" if name == "[scalar]" else "/" + name}{position}'
            continue
        curated[key] = {
            "Dandiset ID": config.dandiset_id,
            "Asset ID": prop.get("asset_id", ""),
            "Audit Field": name,
            "Extracted Value": value,
            "Source Type": prop.get("evidence_kind", "NWB property"),
            "Source Location": f'{prop.get("asset_path", "")}:{path}{"" if name in ("[scalar]", "[array]") else "/" + name}{position}',
            "Review Status": prop.get("read_status", "Read"),
            "Notes": "Value truncated in inventory" if prop.get("value_truncated") else "",
            "Evidence Role": "NWB evidence",
            "Occurrences": 1,
        }
    for rec in curated.values():
        if rec["Occurrences"] > 1:
            rec["Notes"] = (rec["Notes"] + " " if rec["Notes"] else "") + \
                           f'{rec["Occurrences"]} occurrences; all paths and rows remain in NWB Properties.'
        source_rows.append(rec)
    source_rows.extend({**ev.as_dict(config.dandiset_id), "Evidence Role": "NWB mapping candidate"}
                       for ev in result.evidence if ev.source_type in {"NWB alias candidate", "NWB-derived"})
    source_rows.extend(result.dandi_reference_rows)
    return source_rows


def write_master_workbook(config: AuditConfig, result: AuditResult) -> Path:
    _ensure_master(config)
    _backup(config)
    try:
        workbook = load_workbook(config.master_workbook)
    except PermissionError as exc:
        raise PermissionError(f"Close {config.master_workbook.name} in Excel and run again.") from exc

    for sheet in list(OBSOLETE_SHEETS):
        if sheet in workbook.sheetnames:
            del workbook[sheet]

    audit_sheet, metadata_headers, selected = _populate_audit(workbook, config, result)
    _upsert_inventory(workbook, config, result, audit_sheet)

    source_rows = _source_records(config, result)
    requirement_rows = _requirement_review(config, result)
    datasets = {
        "Requirement Review": requirement_rows,
        "NWB Properties": result.nwb_property_rows,
        "DANDI Resources": result.associated_resources,
        "Source Records": source_rows,
        "NWB File Metadata": result.nwb_rows,
        "Dataset Features": _dataset_features(config, result),
        "NWB Findings": _nwb_findings(config, result),
        "NWB Objects": result.nwb_object_rows,
        "Electrode Groups": result.electrode_rows,
        "Asset Inventory": result.asset_rows,
    }
    for sheet, rows in datasets.items():
        _replace_dataset_records(workbook, sheet, rows, CORE_LONG_SHEETS[sheet], config.dandiset_id)

    _write_cde(workbook, config, requirement_rows)
    _write_requirements(workbook)

    # Put the most useful sheets first.
    preferred = [
        "README", "Core Workflow", "Dataset Inventory", "Ephys Audit", "Fiber Photometry Audit", "Calcium Imaging Audit",
        "Other Modalities Audit", "NWB Findings", "Common Data Elements", "Dataset Features", "Source Records",
        "Requirement Review", "DANDI Resources", "Asset Inventory", "NWB File Metadata", "NWB Objects", "Electrode Groups", "NWB Properties", "Allen Requirements", "Field Catalog",
    ]
    ordered = [workbook[name] for name in preferred if name in workbook.sheetnames]
    ordered += [ws for ws in workbook.worksheets if ws.title not in preferred]
    workbook._sheets = ordered

    temp = config.master_workbook.with_name(f".{config.master_workbook.stem}.saving{config.master_workbook.suffix}")
    try:
        workbook.save(temp)
        temp.replace(config.master_workbook)
    except PermissionError as exc:
        temp.unlink(missing_ok=True)
        raise PermissionError(f"Close {config.master_workbook.name} in Excel and run again.") from exc
    return config.master_workbook
